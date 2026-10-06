#!/usr/bin/env python3
"""v1.13.2: series(시리즈/상품 코드) 필드 — Excel ↔ JSON 왕복에서 보존되는지, 옛 Excel(열 없음)과도 호환되는지"""
import json, sys, shutil, tempfile, copy
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import cards_xlsx as X  # noqa
import openpyxl  # noqa

ok_n = bad = 0
def ok(c, m):
    global ok_n, bad
    print(("✓ " if c else "✗ ") + m); ok_n += bool(c); bad += (not c)

db = json.load(open(ROOT / "data" / "cards.json", encoding="utf-8"))
sub_ids = [k for k, v in db["cards"].items() if v.get("series")][:40] + [k for k, v in db["cards"].items() if not v.get("series")][:3]
small = {"cards": {k: db["cards"][k] for k in sub_ids}}
ok(any(v.get("series") for v in small["cards"].values()), f"series 가 있는 카드가 DB 에 있음 ({sum(1 for v in db['cards'].values() if v.get('series'))}장)")
tmp = Path(tempfile.mkdtemp()); xp = tmp / "c.xlsx"
X.write_workbook(small, xp)
wb = openpyxl.load_workbook(xp); ws = wb["Cards"]; hdr = [c.value for c in ws[1]]
ok("series" in hdr and hdr.index("series") == hdr.index("n") + 1, "Cards 시트에 Series 열(카드명 바로 옆)이 있음: " + ",".join(map(str, hdr[:12])))
cards, meta, probs = X.read_workbook(xp)
ok(not [p for p in probs if p.level == "ERROR"], "xlsx 읽기 오류 없음")
merged = X.merge_with_base(cards, meta, small)
ok(all(merged["cards"][k].get("series", "") == small["cards"][k].get("series", "") for k in sub_ids), "Excel → JSON 변환 후 series 보존")
# Excel 에서 직접 수정
col = hdr.index("series") + 1; r = 2; cid = ws.cell(r, 1).value
ws.cell(r, col).value = "P99"; wb.save(xp)
cards, meta, probs = X.read_workbook(xp); merged = X.merge_with_base(cards, meta, small)
ok(merged["cards"][cid]["series"] == "P99", f"Excel 에서 고친 값이 JSON 에 반영됨 ({cid} → P99)")
# 빈 칸으로 지우면 빈 값
ws.cell(r, col).value = None; wb.save(xp)
cards, meta, probs = X.read_workbook(xp); merged = X.merge_with_base(cards, meta, small)
ok(merged["cards"][cid].get("series", "") == "", "Excel 에서 지우면 빈 값")
# 옛 Excel(Series 열 없음): 오류 없이 읽히고 기존 JSON 의 series 유지
wb = openpyxl.load_workbook(xp); ws = wb["Cards"]; ws.delete_cols(hdr.index("series") + 1); wb.save(xp)
cards, meta, probs = X.read_workbook(xp)
ok(not [p for p in probs if p.level == "ERROR"], "Series 열이 없는 옛 Excel 도 오류 없이 읽힘")
merged = X.merge_with_base(cards, meta, small)
ok(all(merged["cards"][k].get("series", "") == small["cards"][k].get("series", "") for k in sub_ids), "옛 Excel 변환 시 기존 JSON series 유지")
# 데이터 형식
bad_fmt = [k for k, v in db["cards"].items() if v.get("series", "") != "" and not __import__("re").fullmatch(r"[PD]\d{2}", v["series"])]
ok(not bad_fmt, f"series 형식(P01~ / D01~) 위반 없음 {bad_fmt[:5]}")
shutil.rmtree(tmp, ignore_errors=True)
print(f"series_xlsx_test: {ok_n} 통과 / {bad} 실패"); sys.exit(1 if bad else 0)
