#!/usr/bin/env python3
"""v1.14.0: 카드 이미지 파일 방식 — Excel(image_file) ↔ JSON(file/img 경로) ↔ 실제 파일, 검증, 새 카드 추가(파일 복사·재압축 없음), 실제 DB 점검.
실행: python3 test/image_files_test.py   (실제 CardImage 폴더가 없으면 '실제 DB 파일 존재' 항목만 건너뜀 — 환경변수 CARD_IMAGE_DIR 로 폴더 지정 가능)"""
import copy, hashlib, json, os, shutil, subprocess, sys, tempfile, warnings
from pathlib import Path
warnings.simplefilter("ignore")
REAL_IMG = os.environ.pop("CARD_IMAGE_DIR", None)   # 임시 프로젝트는 항상 자기 CardImage 폴더를 쓴다 (실제 폴더를 건드리지 않음)
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import cards_xlsx as X  # noqa
import card_images as CI  # noqa
import build_cards_from_excel as B  # noqa
import export_cards_to_excel as EX  # noqa
import openpyxl  # noqa
from PIL import Image  # noqa

ok_n = bad = 0
def ok(c, m):
    global ok_n, bad
    print(("✓ " if c else "✗ ") + m); ok_n += bool(c); bad += (not c)

REAL = json.load(open(ROOT / "data" / "cards.json", encoding="utf-8"))
ids = sorted(REAL["cards"])[:10] + ["id_1170"] * 0
small = {"cards": {k: copy.deepcopy(REAL["cards"][k]) for k in ids}}
tmp = Path(tempfile.mkdtemp()); (tmp / "data").mkdir(); IMG = tmp / "CardImage"; IMG.mkdir()
def mkimg(name, w=700, h=980, color=(200, 30, 30)):
    Image.new("RGB", (w, h), color).save(IMG / name, "JPEG", quality=92)
for k in ids: mkimg(CI_name := f"{k}.jpg", color=(abs(hash(k)) % 255, 80, 120))
mkimg("alt_picture.jpg", 800, 1100, (10, 200, 10)); mkimg("small.jpg", 300, 420)
for k, c in small["cards"].items(): c["file"] = f"{k}.jpg"; c["img"] = f"CardImageWeb/{k}.webp"
js, xp = tmp / "data" / "cards.json", tmp / "data" / "cards.xlsx"
js.write_text(json.dumps(small, ensure_ascii=False), "utf-8")

# ── 1. JSON → Excel: image_file 열
EX.export(js, xp, rebuild=True, quiet=True)
wb = openpyxl.load_workbook(xp); ws = wb["Cards"]; hdr = [c.value for c in ws[1]]
ok("image_file" in hdr and "file" not in hdr, "Cards 시트에 image_file 열이 있음 (옛 'file' 이름은 없음): " + ",".join(map(str, hdr[10:])))
col = hdr.index("image_file") + 1
ok(all(ws.cell(r, col).value == f"{ws.cell(r, 1).value}.jpg" for r in range(2, ws.max_row + 1)), "image_file 값 = 카드ID.jpg")
ok("700x980" in str(ws.cell(2, hdr.index("img_info") + 1).value), "img_info 에 실제 이미지 해상도 표시: " + str(ws.cell(2, hdr.index("img_info") + 1).value))
ok(X.DESC["image_file"][1] == "logic", "image_file 은 게임에 반영되는 열로 표시")

# ── 2. Excel → JSON (round trip: 이미지 연결 정보가 사라지지 않음)
code, info = B.build(xp, js, js, use_node=False, quiet=True)
after = json.loads(js.read_text("utf-8"))
ok(code == 0 and after == small, "JSON → Excel → JSON 왕복 후 모든 카드(이미지 필드 포함)가 그대로")
ok("data:image" not in js.read_text("utf-8"), "cards.json 에 base64 이미지가 없음")

# ── 3. Excel 에서 image_file 고치면 게임(JSON)에 반영 (다른 필드는 불변)
cid = ws.cell(2, 1).value; ws.cell(2, col).value = "alt_picture.jpg"; wb.save(xp)
code, _ = B.build(xp, js, js, use_node=False, quiet=True); new = json.loads(js.read_text("utf-8"))
ok(code == 0 and new["cards"][cid]["file"] == "alt_picture.jpg" and new["cards"][cid]["img"] == "CardImageWeb/alt_picture.webp", f"{cid}: image_file=alt_picture.jpg → img=CardImageWeb/alt_picture.webp")
diff = {k: v for k, v in new["cards"][cid].items() if v != small["cards"][cid].get(k)}
ok(set(diff) == {"file", "img"}, f"바뀐 필드는 file/img 뿐 (AP/LP/color/lv/trait/series/fx/extra/ab 불변): {sorted(diff)}")
ok(all(new["cards"][k] == small["cards"][k] for k in ids if k != cid), "다른 카드는 전혀 바뀌지 않음")

# ── 4. 잘못된 image_file 검증
def try_val(v):
    wb2 = openpyxl.load_workbook(xp); w2 = wb2["Cards"]; w2.cell(3, col).value = v; p = tmp / "t.xlsx"; wb2.save(p)
    cards, meta, probs = X.read_workbook(p); pr = X.validate_cards(cards, small, use_node=False, img_dir=IMG) if cards else probs
    return [q for q in probs + pr if q.level == "ERROR"], [q for q in probs + pr if q.level != "ERROR"]
for bad_v, why in [("../server.js", "상위 폴더 이탈(..)"), ("sub/x.jpg", "하위 경로"), ("C:\\x.jpg", "윈도 경로"), ("http://evil.example/a.jpg", "URL"), ("data:image/jpeg;base64,AAAA", "data URL"),
                   ("x.gif", "허용되지 않는 확장자"), ("noext", "확장자 없음"), ("missing_file.jpg", "폴더에 없는 파일"), ("한글.jpg", "허용 안 되는 문자")]:
    er, _w = try_val(bad_v); ok(any("image_file" in e.where for e in er), f"잘못된 image_file 거부 — {why}: {bad_v!r}")
er, w = try_val(None); ok(not er and any("image_file 이 비어" in x.msg for x in w), "image_file 을 비우면 오류가 아니라 경고(이미지 없는 카드)")
er, w = try_val("alt_picture.jpg"); ok(not er and any("여러 카드가 사용" in x.msg for x in w), "같은 파일을 두 카드가 쓰면 경고")
wb2 = openpyxl.load_workbook(xp); wb2["Cards"].cell(3, col).value = "../etc/passwd.jpg"; bp = tmp / "bad.xlsx"; wb2.save(bp)
before_txt = js.read_text("utf-8"); c2, _i = B.build(bp, js, js, use_node=False, quiet=True)
ok(c2 != 0 and js.read_text("utf-8") == before_txt, "잘못된 image_file 이 있으면 빌드 중단(cards.json 은 변경되지 않음)")
# 이미지 파일 손상
(IMG / "broken.jpg").write_bytes(b"not an image at all")
wb2 = openpyxl.load_workbook(xp); wb2["Cards"].cell(3, col).value = "broken.jpg"; wb2.save(bp); cards, meta, probs = X.read_workbook(bp)
pr = X.validate_cards(cards, small, use_node=False, img_dir=IMG); ok(any("읽을 수 없습니다" in q.msg for q in pr), "깨진 이미지 파일은 오류로 보고")
# 저해상도 경고
wb2 = openpyxl.load_workbook(xp); wb2["Cards"].cell(3, col).value = "small.jpg"; wb2.save(bp); cards, meta, probs = X.read_workbook(bp)
pr = X.validate_cards(cards, small, use_node=False, img_dir=IMG); ok(any("해상도가 낮" in q.msg and q.level == "WARN" for q in pr), "저해상도(폭 300px) 이미지는 경고")
# 폴더에 있으나 연결 안 된 파일 안내
ok(any("연결되지 않은 파일" in q.msg for q in pr), "어떤 카드와도 연결되지 않은 이미지 파일을 안내")

# ── 5. 옛 Excel(file 열 이름) 도 읽힘
wb3 = openpyxl.load_workbook(xp); w3 = wb3["Cards"]; w3.cell(1, col).value = "file"; w3.cell(2, col).value = f"{cid}.jpg"; w3.cell(3, col).value = f"{ws.cell(3, 1).value}.jpg"; op = tmp / "old.xlsx"; wb3.save(op)
cards, meta, probs = X.read_workbook(op); ok(not [q for q in probs if q.level == "ERROR"] and cards[cid]["file"] == f"{cid}.jpg", "옛 Excel(열 이름 file)도 오류 없이 읽힘 → file 값 사용")
# 열이 없는 아주 옛 Excel: 기존 JSON 값 유지
wb3 = openpyxl.load_workbook(xp); w3 = wb3["Cards"]; w3.delete_cols(col); op2 = tmp / "old2.xlsx"; wb3.save(op2)
cards, meta, probs = X.read_workbook(op2); mg = X.merge_with_base(cards, meta, small)
ok(all(mg["cards"][k]["file"] == small["cards"][k]["file"] and mg["cards"][k]["img"] == small["cards"][k]["img"] for k in ids), "image_file 열이 없는 Excel → 기존 JSON 의 이미지 연결 유지")
# 옛 base64 가 JSON 에 남아 있어도 빌드하면 경로로 바뀜 (base64 제거)
leg = copy.deepcopy(small)
for k, c in leg["cards"].items(): c["img"] = "data:image/jpeg;base64,/9j/4AAQ"; c["file"] = f"ID_{k}_원본이름.jpg"
cards, meta, probs = X.read_workbook(xp); mg = X.merge_with_base(cards, meta, leg)
ok(all(mg["cards"][k]["img"].startswith("CardImageWeb/") and "data:" not in mg["cards"][k]["img"] for k in ids), "옛 base64 가 남은 JSON 에서도 Excel 의 image_file 기준으로 경로가 만들어짐(base64 제거)")

# ── 6. 새 카드 추가: 원본 파일 그대로 복사(재압축 없음) → Excel 에 image_file → JSON 에는 경로만
EX.export(js, xp, rebuild=True, quiet=True); small2 = json.loads(js.read_text("utf-8"))
src = tmp / "new"; src.mkdir(); big = src / "ID_9500_テスト_R.jpg"; Image.new("RGB", (716, 1000), (9, 99, 199)).save(big, "JPEG", quality=95)
src_bytes = big.read_bytes()
tj = tmp / "tr.json"; tj.write_text(json.dumps([{"id": "9500", "n": "テスト", "type": "char", "color": "blue", "lv": "1", "ap": "1000", "lp": "1", "trait": "", "kw": "", "fx": "", "extra": "", "src": big.name}], ensure_ascii=False), "utf-8")
r = subprocess.run([sys.executable, str(ROOT / "add_new_cards.py"), str(src), "--from-json", str(tj), "--db", str(js)], capture_output=True, text=True, cwd=str(tmp), timeout=300)
out = r.stdout + r.stderr; nj = json.loads(js.read_text("utf-8")); nc = nj["cards"].get("id_9500")
ok(r.returncode == 0 and nc is not None, "add_new_cards --from-json 로 새 카드 추가: " + (out[-200:] if r.returncode else "OK"))
ok(nc and nc["file"] == "id_9500.jpg" and nc["img"] == "CardImageWeb/id_9500.webp", "새 카드: file=id_9500.jpg, img=CardImageWeb/id_9500.webp (경로만 기록)")
ok("data:image" not in js.read_text("utf-8"), "새 카드를 추가해도 cards.json 에 base64 가 들어가지 않음")
ok((IMG / "id_9500.jpg").exists() and (IMG / "id_9500.jpg").read_bytes() == src_bytes, "이미지는 CardImage/id_9500.jpg 로 원본 그대로 복사(바이트 동일, 재압축/축소 없음)")
ok(CI.read_size(IMG / "id_9500.jpg") == (716, 1000), "해상도 716x1000 유지")
ok(all(nj["cards"][k] == small2["cards"][k] for k in small2["cards"]), "기존 카드 10장은 전혀 바뀌지 않음")
wb4 = openpyxl.load_workbook(xp); w4 = wb4["Cards"]; h4 = [c.value for c in w4[1]]; rows = {w4.cell(r, 1).value: r for r in range(2, w4.max_row + 1)}
ok("id_9500" in rows and w4.cell(rows["id_9500"], h4.index("image_file") + 1).value == "id_9500.jpg", "cards.xlsx 에 새 카드 행 + image_file=id_9500.jpg 기록")
code, _ = B.build(xp, js, js, use_node=False, quiet=True); ok(code == 0 and json.loads(js.read_text("utf-8")) == nj, "추가 후 Excel → JSON 빌드가 동일 결과 (왕복 일치)")
# 이미지 없는 신탄(원본 파일이 없는 경우) 은 추가되지 않고 안전하게 실패
tj2 = tmp / "tr2.json"; tj2.write_text(json.dumps([{"id": "9501", "n": "없음", "type": "char", "color": "blue", "lv": "1", "ap": "1000", "lp": "1", "src": "nofile.jpg"}], ensure_ascii=False), "utf-8")
b4 = js.read_text("utf-8"); r = subprocess.run([sys.executable, str(ROOT / "add_new_cards.py"), str(src), "--from-json", str(tj2), "--db", str(js)], capture_output=True, text=True, cwd=str(tmp), timeout=300)
ok(js.read_text("utf-8") == b4 and not (IMG / "id_9501.jpg").exists(), "원본 이미지를 찾지 못하면 카드/파일 모두 추가되지 않음")

# ── 7. 헬퍼 단위
ok(CI.check_name("id_1068.jpg") is None and CI.check_name("a/b.jpg") and CI.check_name("..\\a.jpg") and CI.check_name("a.bmp") and CI.check_name(" a.jpg"), "check_name: 정상/경로/확장자/공백 판정")
ok(CI.img_path("id_1.jpg") == "CardImageWeb/id_1.webp" and CI.img_path("") == "", "img_path")
shutil.rmtree(tmp, ignore_errors=True)

# ── 8. 실제 프로젝트 데이터
db = REAL["cards"]; cs, meta, probs = X.read_workbook(ROOT / "data" / "cards.xlsx")
ok(not [p for p in probs if p.level == "ERROR"], "실제 data/cards.xlsx 읽기 오류 없음")
ok(len(db) >= 1339 and all("data:image" not in (c.get("img") or "") for c in db.values()), f"실제 cards.json {len(db)}장: base64 이미지 0건")
ok(all(c.get("img") == CI.img_path(c.get("file")) and c.get("file") for c in db.values()), "모든 카드: img == CardImageWeb/<file 의 .webp>, file 이 비어 있지 않음")
ok(all(CI.check_name(c["file"]) is None for c in db.values()), "모든 file 이 안전한 파일명 형식")
ok(all(cs[k].get("file") == db[k].get("file") for k in db) and set(cs) == set(db), "Excel image_file 과 JSON file 이 전 카드에서 일치")
ok(all(db[k]["file"].rsplit(".", 1)[0] == k for k in db), "파일명이 모두 카드 ID 와 같음 (id_XXXX.jpg)")
rd = Path(REAL_IMG) if REAL_IMG else ROOT / "CardImage"
if rd.is_dir():
    miss = [k for k, c in db.items() if not (rd / c["file"]).is_file()]
    ok(not miss, f"실제 CardImage: 모든 카드의 이미지 파일 존재 (누락 {len(miss)}) {miss[:5]}")
    sz = [CI.read_size(rd / c["file"]) for c in db.values() if (rd / c["file"]).is_file()]
    ok(sz and all(s and min(s) >= 470 for s in sz), f"실제 이미지 해상도: 전부 짧은 변 470px 이상 (최소 {min(min(s) for s in sz if s)}px, 이전 base64 는 172x240)")
else: print(f"- 건너뜀: 실제 CardImage 폴더가 없음 ({rd})")
print(f"image_files_test: {ok_n} 통과 / {bad} 실패"); sys.exit(1 if bad else 0)
